import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_MD = ROOT / "docs" / "project-report.md"
AI_LOG_MD = ROOT / "docs" / "ai-usage-log.md"
TEST_CASES_MD = ROOT / "docs" / "test-cases.md"
OUTPUT = ROOT / "docs" / "MongoForm-第一阶段项目报告.docx"

BODY_FONT = "Noto Sans SC"
HEADING_FONT = "Noto Serif SC"
MONO_FONT = "Consolas"


def set_run_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=85, start=90, bottom=85, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, BODY_FONT, 8.5, color=(90, 90, 90))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, BODY_FONT, 8.5, color=(90, 90, 90))


def configure_document(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.space_before = Pt(0)

    for style_name, size, before, after in (
        ("Title", 30, 0, 16),
        ("Heading 1", 20, 14, 8),
        ("Heading 2", 14, 12, 6),
        ("Heading 3", 11.5, 9, 5),
    ):
        style = document.styles[style_name]
        style.font.name = HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if style_name == "Title":
            p_pr = style._element.get_or_add_pPr()
            borders = p_pr.find(qn("w:pBdr"))
            if borders is not None:
                p_pr.remove(borders)

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.3

    add_page_number(section.footer.paragraphs[0])
    document.core_properties.title = "基于 MongoDB 的在线问卷系统第一阶段项目完成报告"
    document.core_properties.subject = "在线问卷系统设计、实现、测试与 AI 使用说明"
    document.core_properties.author = "项目开发者"
    document.core_properties.keywords = "MongoDB, Flask, 问卷系统, 条件跳转, AI"


def add_cover(document):
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)

    eyebrow = document.add_paragraph()
    run = eyebrow.add_run("第一阶段项目报告")
    set_run_font(run, BODY_FONT, 10, True, (80, 80, 80))
    eyebrow.paragraph_format.space_after = Pt(22)

    title = document.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    run = title.add_run("基于 MongoDB 的在线问卷系统")
    set_run_font(run, HEADING_FONT, 29, True, (0, 0, 0))
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    subtitle = document.add_paragraph()
    run = subtitle.add_run("项目设计、实现、测试与 AI 使用说明")
    set_run_font(run, BODY_FONT, 14, False, (0, 0, 0))
    subtitle.paragraph_format.space_after = Pt(85)

    info = document.add_table(rows=5, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.LEFT
    info.autofit = False
    values = [
        ("技术栈", "Python / Flask / MongoDB / 原生前端"),
        ("数据库", "MongoDB 8.3，五个业务集合"),
        ("测试结果", "13 项自动测试通过"),
        ("日期", "2026 年 10 月 22 日"),
    ]
    for index, (cell, text) in enumerate(
        zip(info.rows[0].cells, ("项目信息", "内容"))
    ):
        cell.width = Inches(1.35 if index == 0 else 4.75)
        set_cell_shading(cell, "244B5A")
        set_cell_border(cell)
        set_cell_margins(cell, 110, 110, 110, 110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, BODY_FONT, 9, True, (255, 255, 255))
    set_repeat_table_header(info.rows[0])

    for row, (label, value) in zip(info.rows[1:], values):
        row.cells[0].width = Inches(1.35)
        row.cells[1].width = Inches(4.75)
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell, 110, 110, 110, 110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        set_run_font(r0, BODY_FONT, 9, True, (0, 0, 0))
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(value)
        set_run_font(r1, BODY_FONT, 9, False, (0, 0, 0))

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(50)
    run = note.add_run("项目开发者")
    set_run_font(run, BODY_FONT, 10, True, (0, 0, 0))
    document.add_page_break()


def add_inline(paragraph, value, base_size=10.5):
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    cursor = 0
    for match in pattern.finditer(value):
        if match.start() > cursor:
            run = paragraph.add_run(value[cursor : match.start()])
            set_run_font(run, BODY_FONT, base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, BODY_FONT, base_size, True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, MONO_FONT, base_size - 0.8)
        cursor = match.end()
    if cursor < len(value):
        run = paragraph.add_run(value[cursor:])
        set_run_font(run, BODY_FONT, base_size)


def column_widths(count):
    maps = {
        1: [6.5],
        2: [1.7, 4.8],
        3: [1.25, 2.45, 2.8],
        4: [1.0, 1.85, 1.75, 1.9],
        5: [0.72, 1.8, 1.6, 1.7, 0.68],
        6: [0.65, 1.25, 1.15, 1.1, 1.7, 0.65],
    }
    if count in maps:
        return maps[count]
    return [6.5 / count] * count


def add_markdown_table(document, rows):
    if not rows:
        return
    headers = [cell.strip() for cell in rows[0].strip("|").split("|")]
    body = []
    for row in rows[1:]:
        if re.match(r"^\s*\|?\s*:?-{3,}", row):
            continue
        values = [cell.strip() for cell in row.strip("|").split("|")]
        values += [""] * (len(headers) - len(values))
        body.append(values[: len(headers)])

    table = document.add_table(rows=1 + len(body), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = column_widths(len(headers))

    for col_index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[col_index])
        set_cell_shading(cell, "244B5A")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(header)
        set_run_font(run, BODY_FONT, 8.2, True, (255, 255, 255))

    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for row_index, values in enumerate(body, start=1):
        prevent_row_split(table.rows[row_index])
        for col_index, value in enumerate(values):
            cell = table.rows[row_index].cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, "F4F6F7")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if len(body) <= 6 and row_index < len(body):
                p.paragraph_format.keep_with_next = True
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if col_index == 0 or len(value) <= 8
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline(p, value, 8.2)

    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(document, lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.25)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F3F4")
    p_pr.append(shading)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, MONO_FONT, 8.2, color=(35, 35, 35))


def markdown_to_docx(document, markdown, base_heading=1, chapter_breaks=True):
    lines = markdown.splitlines()
    index = 0
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        value = " ".join(part.strip() for part in paragraph_buffer)
        paragraph_buffer = []
        p = document.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, value)

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            if document.paragraphs:
                document.paragraphs[-1].paragraph_format.keep_with_next = True
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(lines[index])
                index += 1
            add_markdown_table(document, table_rows)
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading_match:
            flush_paragraph()
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            if base_heading == 1 and level == 1:
                index += 1
                continue
            mapped = min(3, max(1, level - base_heading + 1))
            p = document.add_paragraph(style=f"Heading {mapped}")
            if (
                chapter_breaks
                and level == 2
                and re.match(r"^\d+\.", text)
                and not text.startswith("1. 项目")
            ):
                p.paragraph_format.page_break_before = True
            run = p.add_run(text)
            set_run_font(run, HEADING_FONT, 20 if mapped == 1 else 14 if mapped == 2 else 11.5, True, (0, 0, 0))
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if bullet or numbered:
            flush_paragraph()
            if bullet:
                p = document.add_paragraph(style="List Bullet")
                add_inline(p, bullet.group(1))
            else:
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.75)
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, f"{numbered.group(1)}. {numbered.group(2)}")
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.right_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            add_inline(p, stripped.lstrip("> ").strip())
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()


def main():
    document = Document()
    configure_document(document)
    add_cover(document)

    report = REPORT_MD.read_text(encoding="utf-8")
    report_body = report[report.index("## 摘要") :]
    markdown_to_docx(document, report_body, base_heading=2, chapter_breaks=False)

    appendix_heading = document.add_paragraph(style="Heading 1")
    appendix_heading.paragraph_format.page_break_before = True
    run = appendix_heading.add_run("附录 A AI 使用日志")
    set_run_font(run, HEADING_FONT, 20, True, (0, 0, 0))
    ai_log = AI_LOG_MD.read_text(encoding="utf-8")
    ai_body = ai_log[ai_log.index("## 1. 记录范围") :]
    markdown_to_docx(document, ai_body, base_heading=2, chapter_breaks=False)

    appendix_heading = document.add_paragraph(style="Heading 1")
    appendix_heading.paragraph_format.page_break_before = True
    run = appendix_heading.add_run("附录 B 测试用例与结果")
    set_run_font(run, HEADING_FONT, 20, True, (0, 0, 0))
    test_cases = TEST_CASES_MD.read_text(encoding="utf-8")
    test_body = test_cases[test_cases.index("## 1. 测试范围") :]
    markdown_to_docx(document, test_body, base_heading=2, chapter_breaks=False)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
